import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_bg(cell, col):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{col}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(0.7)
    s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("TeamFlow — Tizimga Kirish Login va Parollari")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(30, 58, 138)

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_info = info.add_run("Kirish manzili: http://192.168.224.70:5183/login  |  Admin: http://192.168.224.70:8010/admin/")
r_info.font.size = Pt(10)
r_info.font.bold = True
r_info.font.color.rgb = RGBColor(2, 132, 199)

p_sep = doc.add_paragraph()
p_sep.paragraph_format.space_after = Pt(8)

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Ism / Xodim", "Lavozim", "Rol", "Login (Email)", "Parol"]

for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    run = cell.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(255, 255, 255)
    set_bg(cell, "1E3A8A")

users = [
    ("Bosh Admin", "Tizim Administratori", "ADMIN", "admin@teamflow.uz", "admin12345"),
    ("Akmal Boshliqov", "Kompaniya Rahbari / Boshliq", "BOSS", "boshliq@teamflow.uz", "password123"),
    ("Sardor Menejerov", "Loyiha Menejeri (Lead PM)", "MANAGER", "pm@teamflow.uz", "password123"),
    ("Bobur Rahimov", "Loyiha Menejeri", "MANAGER", "menejer@teamflow.uz", "menejer123"),
    ("Jasur Backendchi", "Senior Backend Dasturchi", "DEVELOPER", "backend@teamflow.uz", "password123"),
    ("Malika Frontendchi", "Middle Frontend Dasturchi", "DEVELOPER", "frontend@teamflow.uz", "password123"),
    ("Dilshod Karimov", "Tizim Operatori", "OPERATOR", "operator@teamflow.uz", "operator123"),
]

for idx, (name, title, role, email, pwd) in enumerate(users):
    row = table.add_row().cells
    bg = "F1F5F9" if idx % 2 == 1 else "FFFFFF"
    for c_idx, val in enumerate([name, title, role, email, pwd]):
        row[c_idx].text = val
        run = row[c_idx].paragraphs[0].runs[0]
        run.font.size = Pt(9.5)
        if c_idx == 0:
            run.font.bold = True
        elif c_idx == 2:
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
        elif c_idx == 4:
            run.font.bold = True
            run.font.color.rgb = RGBColor(22, 101, 52)
        set_bg(row[c_idx], bg)

doc.save('/app/TeamFlow_Login_Parollar.docx')
print('OK')
